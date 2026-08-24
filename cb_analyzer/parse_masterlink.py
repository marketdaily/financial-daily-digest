"""元富證券『CB初級市場資訊』Excel 解析器(老闆轉寄來源)。

與元大『CB發行案件更新』欄位排列不同,但輸出【完全相同】的正規化 schema,
下游 cb.py 評分 / cb_server / 網站零改動。由 parse_excel.parse() 自動偵測格式後委派。

兩段結構:
  「送件及詢圈標的」(已在市場、完整定價)  表頭 A標的代號 B名稱 C TCRI D承銷方式 E金額(億)
      F發行期間 G發行價格 H承銷券商(擔保銀) I溢價率 J轉換價 K轉換價值 L年/賣回價
      M收件日 N生效日 O詢圈競拍期間 P掛牌日 Q CB拆解日 R資本額 S股價 T CB價位 U備註
  「董事會公告發行標的」(剛通過、未定價)  I 欄改為『董事會通過日』,無轉換價/溢價率
"""
import os
import re
import json
import datetime
import openpyxl
from parse_excel import (_s, _num, _parse_premium, _parse_split_date, DB_PATH)

SEC1 = "送件及詢圈標的"
SEC2 = "董事會公告發行標的"


def _ml_underwriter(raw):
    """'華南永昌證(土地銀)' → ('華南永昌證','土地銀',True);'凱基證' → ('凱基證','',False)。"""
    raw = (raw or "").strip()
    m = re.search(r"[（(]([^）)]*)[）)]", raw)
    coll = m.group(1).strip() if m else ""
    und = re.sub(r"[（(][^）)]*[）)]", "", raw).strip()
    secured = bool(coll) and coll not in ("無擔", "無擔保")
    return und, coll, secured


def _ml_put(raw, tenor):
    """'年/賣回價':'3年100' / '2年101.0025' / '2年100~101.0025' / '2年101~10302'(髒)。
    元富直接給賣回價 → 塞 redeem_price(redemption_value 優先採用,零推估)。"""
    raw = (raw or "").replace("～", "~").strip()
    m = re.search(r"(\d+)\s*年", raw)
    year = int(m.group(1)) if m else tenor
    rest = raw.split("年", 1)[1] if "年" in raw else raw
    nums = [_num(x) for x in re.findall(r"\d+(?:\.\d+)?", rest)]
    nums = [n for n in nums if n is not None and 90 <= n <= 130]  # 合理賣回價區間,濾髒值
    redeem = max(nums) if nums else 100.0
    return {"kind": "YTP", "year": year, "redeem_price": redeem,
            "y_low": None, "y_high": None, "y_mid": None, "raw": raw}


def _ml_auction(note):
    low = high = None
    m = re.search(r"最低標[：:\s]*\s*(\d{2,3}(?:\.\d+)?)", note or "")
    if m:
        low = _num(m.group(1))
    m = re.search(r"最高標[：:\s]*\s*(\d{2,3}(?:\.\d+)?)", note or "")
    if m:
        high = _num(m.group(1))
    return low, high


def _ml_clearing(issue_px_raw, bookbuild, note):
    """發行價格 G:競拍=得標承銷價、詢圈=固定發行價。'不低於100' = 競拍未定(底標)。"""
    raw = (issue_px_raw or "").strip()
    low, high = _ml_auction(note)
    mm = re.search(r"(\d{2,3}(?:\.\d+)?)", raw)
    num = _num(mm.group(1)) if mm else None
    floor_only = "不低於" in raw
    if bookbuild == "競拍":
        if num is not None and not floor_only:
            clearing, method = num, "競拍承銷價"
        elif low and high:
            clearing, method = (low + high) / 2, "競拍中值"
        elif num is not None and floor_only:
            clearing, method = num, "競拍未定(底標)"
        else:
            clearing, method = None, "未定"
    else:  # 詢圈/固定
        clearing, method = (num, "固定發行") if num is not None else (None, "未定")
    return {"clearing": clearing, "auction_low": low, "auction_high": high, "method": method}


def _row1(r, section):
    """section 1『送件及詢圈標的』完整定價列。"""
    code = _s(r[0])
    name = _s(r[1])
    if not code or not name or not code[:2].isdigit():
        return None
    try:
        tcri = int(re.search(r"\d+", _s(r[2])).group())
    except Exception:
        tcri = None
    bookbuild = _s(r[3])
    und, coll, secured = _ml_underwriter(_s(r[7]))
    p_lo, p_hi, p_mid = _parse_premium(_s(r[8]))
    tenor = _num(_s(r[5]))
    tenor = int(tenor) if tenor else None
    put = _ml_put(_s(r[11]), tenor)
    note = _s(r[20]) if len(r) > 20 else ""
    clr = _ml_clearing(_s(r[6]), bookbuild, note)
    return {
        "section": section,
        "stock_code": code[:4],
        "bond_code": code,
        "name": name,
        "tcri": tcri,
        "collateral": coll,
        "secured": secured,
        "size_yi": _num(_s(r[4])),
        "underwriter": und,
        "announce_date": _s(r[12]) or None,      # 收件日
        "effective_date": _s(r[13]) or None,     # 生效日
        "bookbuild": bookbuild,
        "premium_low": p_lo, "premium_high": p_hi, "premium_mid": p_mid,
        "conv_price": _num(_s(r[9])),            # 轉換價
        "listing_date": _s(r[15]) or None,       # 掛牌日
        "split_date": _parse_split_date(_s(r[16])),   # CB拆解日
        "put": put,
        "tenor_year": tenor,
        "clearing_price": clr["clearing"],
        "auction_low": clr["auction_low"],
        "auction_high": clr["auction_high"],
        "pricing_method": clr["method"],
        "issue_price": clr["clearing"],
        "note": note,
        "broker_spot": _num(_s(r[18])) if len(r) > 18 else None,   # 券商表列股價(參考,不進定價)
        "broker_cb_price": _num(_s(r[19])) if len(r) > 19 else None,
    }


def _row2(r, section):
    """section 2『董事會公告發行標的』未定價列(I=董事會通過日,無轉換價/溢價率)。"""
    code = _s(r[0])
    name = _s(r[1])
    if not code or not name or not code[:2].isdigit():
        return None
    try:
        tcri = int(re.search(r"\d+", _s(r[2])).group())
    except Exception:
        tcri = None
    bookbuild = _s(r[3])
    und, coll, secured = _ml_underwriter(_s(r[7]))
    tenor = _num(_s(r[5]))
    tenor = int(tenor) if tenor else None
    note = _s(r[20]) if len(r) > 20 else ""
    clr = _ml_clearing(_s(r[6]), bookbuild, note)
    return {
        "section": section,
        "stock_code": code[:4],
        "bond_code": code,
        "name": name,
        "tcri": tcri,
        "collateral": coll,
        "secured": secured,
        "size_yi": _num(_s(r[4])),
        "underwriter": und,
        "announce_date": _s(r[8]) or None,       # 董事會通過日
        "effective_date": None,
        "bookbuild": bookbuild,
        "premium_low": None, "premium_high": None, "premium_mid": None,
        "conv_price": None,
        "listing_date": None,
        "split_date": None,
        "put": _ml_put("", tenor),
        "tenor_year": tenor,
        "clearing_price": clr["clearing"],
        "auction_low": clr["auction_low"],
        "auction_high": clr["auction_high"],
        "pricing_method": clr["method"],
        "issue_price": clr["clearing"],
        "note": note,
        "broker_spot": _num(_s(r[18])) if len(r) > 18 else None,
        "broker_cb_price": _num(_s(r[19])) if len(r) > 19 else None,
    }


def is_masterlink(ws):
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        for c in row[:3]:
            s = _s(c)
            if SEC1 in s or "承銷\n方式" in s or (s == "標的\n代號"):
                return True
        if i >= 6:
            break
    return False


def parse_ws(ws):
    section = ""
    out = []
    for row in ws.iter_rows(values_only=True):
        a = _s(row[0])
        if SEC1 in a:
            section = SEC1
            continue
        if SEC2 in a:
            section = SEC2
            continue
        if a in ("標的\n代號", "代碼") or _s(row[1]) in ("標的\n名稱", "標的名稱"):
            continue  # 表頭
        rec = _row1(row, section) if section == SEC1 else _row2(row, section)
        if rec:
            out.append(rec)
    return out


def parse(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    return parse_ws(wb[wb.sheetnames[0]])


if __name__ == "__main__":
    import sys
    data = parse(sys.argv[1])
    print(f"元富解析 {len(data)} 檔")
    for d in data[:3]:
        print(" ", d["bond_code"], d["name"], "TCRI", d["tcri"],
              "size", d["size_yi"], "conv", d["conv_price"], "clr", d["clearing_price"])
